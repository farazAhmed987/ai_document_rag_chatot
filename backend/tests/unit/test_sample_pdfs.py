"""Unit tests for the Phase 1 document-generation scripts.

SRD Sections 8.1, 12, 15 and Appendix A.
Requirements: DM-2, IN-1 and IN-2.

Tests generate their own tiny fixtures and do not modify
permanent sample data or use external databases.
"""

from __future__ import annotations

import hashlib
import subprocess
import sys
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from app.core.config import REPO_ROOT

STEMS = (
    "Company_Overview",
    "Property_Listings",
    "Services_and_Fees",
    "FAQs",
    "Policies_and_Terms",
)

SCRIPT = REPO_ROOT / "scripts" / "build_sample_pdfs.py"


@pytest.fixture
def markdown_sources(tmp_path: Path) -> Path:
    """Generate five short Markdown files in a temporary folder."""
    source_dir = tmp_path / "sources"
    source_dir.mkdir()

    for stem in STEMS:
        title = stem.replace("_", " ")

        markdown = f"""# {title}

Document title: {title}
Document code: TEST-001
Version: 1.0
Effective date: 2026-10-28
Owner department: Testing

Fictional company created for software testing.

## Overview

This is a **bold** and *italic* paragraph describing
the fictional sample company.

- First bullet item
- Second bullet item

1. First numbered item
2. Second numbered item

## Example Table

| Service | Fee | Description |
| --- | --- | --- |
| Consultation | PKR 500 | Basic consultation |
| Document Review | PKR 1,000 | Full document review |

---

Final paragraph for extraction testing.
"""

        (source_dir / f"{stem}.md").write_text(
            markdown,
            encoding="utf-8",
        )

    return source_dir


def run_builder(source_dir: Path, output_dir: Path) -> None:
    """Run the real build script as an external command."""
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--source-dir",
            str(source_dir),
            "--output-dir",
            str(output_dir),
        ],
        cwd=REPO_ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + "\n" + result.stderr


def test_sample_pdf_has_real_text_and_tables(
    markdown_sources: Path,
    tmp_path: Path,
) -> None:
    """Verify selectable text, headings, tables, header and footer."""
    output_dir = tmp_path / "generated"
    run_builder(markdown_sources, output_dir)

    pdf_path = output_dir / "Services_and_Fees.pdf"
    assert pdf_path.is_file()

    with pymupdf.open(pdf_path) as document:
        assert len(document) >= 1

        first_page = document[0]
        text = first_page.get_text()

        assert "Himalaya Crest Realty" in text
        assert "Services and Fees" in text
        assert "Consultation" in text
        assert "PKR 500" in text
        assert "Final paragraph" in text
        assert "Page 1 of 1" in text

        tables = first_page.find_tables()
        assert len(tables.tables) >= 1


def test_pdfs_are_byte_for_byte_deterministic(
    markdown_sources: Path,
    tmp_path: Path,
) -> None:
    """Generate identical PDFs in independent output directories."""
    first = tmp_path / "first"
    second = tmp_path / "second"

    run_builder(markdown_sources, first)
    run_builder(markdown_sources, second)

    for stem in STEMS:
        first_bytes = (first / f"{stem}.pdf").read_bytes()
        second_bytes = (second / f"{stem}.pdf").read_bytes()

        assert first_bytes == second_bytes
        assert hashlib.sha256(first_bytes).hexdigest() == (
            hashlib.sha256(second_bytes).hexdigest()
        )


def test_variants_preserve_source_content(
    markdown_sources: Path,
    tmp_path: Path,
) -> None:
    """Ensure generated document variants retain source content."""
    output_dir = tmp_path / "generated"
    run_builder(markdown_sources, output_dir)

    variants = output_dir / "variants"
    markdown_source = (markdown_sources / "FAQs.md").read_bytes()

    assert (variants / "FAQs.md").read_bytes() == markdown_source

    txt_content = (variants / "FAQs.txt").read_text(encoding="utf-8")
    assert "Basic consultation" in txt_content
    assert "Fictional company" in txt_content

    docx_path = variants / "Services_and_Fees.docx"
    assert docx_path.exists()

    document = Document(docx_path)
    assert len(document.tables) >= 1

    headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]

    assert headings
    assert any(
        cell.text == "Consultation"
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )


def test_pdf_rejects_non_winansi_characters(
    markdown_sources: Path,
    tmp_path: Path,
) -> None:
    """Fail explicitly instead of corrupting unsupported PDF text."""
    source = markdown_sources / "FAQs.md"

    source.write_text(
        source.read_text(encoding="utf-8") + "\nUnsupported emoji: 🧪\n",
        encoding="utf-8",
    )

    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--source-dir",
            str(markdown_sources),
            "--output-dir",
            str(tmp_path / "generated"),
        ],
        cwd=REPO_ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode != 0
    assert "Windows-1252" in result.stderr
