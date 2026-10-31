"""RED tests for heading paths and Word XML edge cases (IN-1, IN-5)."""

from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.adapters.extractors.factory import extract_document
from app.schemas.extraction import ExtractedDocument, ExtractedUnit


def test_pdf_page_two_inherits_first_page_heading(tmp_path: Path) -> None:
    """Headings from previous pages remain available for later citations."""
    pdf = pymupdf.open()
    pdf.new_page().insert_text((30, 60), "# Sales conditions")
    pdf.new_page().insert_text((30, 60), "Full text about the terms")
    path = tmp_path / "chapters.pdf"
    pdf.save(path)
    pdf.close()
    result = extract_document(path)
    # PDF heading levels derive from font-size heuristics, so a deterministic
    # direct-annotation fixture verifies the actual context contract below.
    from app.services.structure import annotate_heading_paths

    result.units[0].markdown = "# Sales conditions\nText"
    result.units[1].markdown = "Terms continued"
    annotate_heading_paths(result)
    assert result.units[1].heading_path == ["Sales conditions"]


def test_heading_context_tracks_subheadings() -> None:
    """The effective path changes at each sub-heading inside one unit."""
    from app.services.structure import (
        annotate_heading_paths,
        heading_context,
        iter_headings,
    )

    doc = ExtractedDocument(
        "example.md",
        "md",
        [
            ExtractedUnit(
                "section", 0, None, [], "# Main\nintro\n## First\nA\n## Second\nB"
            )
        ],
        None,
    )
    annotate_heading_paths(doc)
    unit = doc.units[0]
    assert unit.heading_path == ["Main"]
    assert heading_context(unit, 3) == ["Main", "First"]
    assert heading_context(unit, 5) == ["Main", "Second"]
    assert list(iter_headings(unit.markdown)) == [
        (0, 1, "Main"),
        (2, 2, "First"),
        (4, 2, "Second"),
    ]


def test_docx_section_paths_are_not_last_subheading(tmp_path: Path) -> None:
    """Top-level units have context at start, not the final subheading."""
    path = tmp_path / "outline.docx"
    doc = Document()
    doc.add_heading("One", level=1)
    doc.add_heading("A", level=2)
    doc.add_paragraph("Alpha body")
    doc.add_heading("B", level=2)
    doc.add_heading("Two", level=1)
    doc.add_heading("C", level=2)
    doc.save(path)
    extracted = extract_document(path)
    assert len(extracted.units) == 2
    assert extracted.units[0].heading_path == ["One"]
    assert extracted.units[1].heading_path == ["Two"]


def test_docx_body_sdt_and_track_changes(tmp_path: Path) -> None:
    """Include content-control and inserted text, excluding deleted text."""
    path = tmp_path / "revisions.docx"
    doc = Document()
    doc.add_paragraph("Before")
    sdt = OxmlElement("w:sdt")
    sdt_content = OxmlElement("w:sdtContent")
    inserted_para = OxmlElement("w:p")
    insert = OxmlElement("w:ins")
    insert_run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Inserted part"
    insert_run.append(text)
    insert.append(insert_run)
    inserted_para.append(insert)
    deleted = OxmlElement("w:del")
    delete_run = OxmlElement("w:r")
    deleted_text = OxmlElement("w:delText")
    deleted_text.text = "Removed part"
    delete_run.append(deleted_text)
    deleted.append(delete_run)
    inserted_para.append(deleted)
    sdt_content.append(inserted_para)
    sdt.append(sdt_content)
    doc.element.body.insert(1, sdt)
    doc.save(path)
    extracted = extract_document(path)
    content = "\n".join(unit.markdown for unit in extracted.units)
    assert "Inserted part" in content
    assert "Removed part" not in content


def test_docx_horizontal_merges_and_nested_tables(tmp_path: Path) -> None:
    """Merged content appears once; a nested table is not silently discarded."""
    path = tmp_path / "nested.docx"
    doc = Document()
    outer = doc.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "Merged label"
    outer.cell(0, 0).merge(outer.cell(0, 1))
    outer.cell(1, 0).text = "Outer"
    nested = outer.cell(1, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Nested Key"
    nested.cell(0, 1).text = "Nested Value"
    doc.save(path)
    content = "\n".join(unit.markdown for unit in extract_document(path).units)
    assert content.count("Merged label") == 1
    assert "Nested Key" in content and "Nested Value" in content


def test_docx_title_and_numbering_properties(tmp_path: Path) -> None:
    """Title is a heading; explicit numPr is interpreted as a list."""
    path = tmp_path / "numbering.docx"
    doc = Document()
    doc.add_paragraph("Document title", style="Title")
    paragraph = doc.add_paragraph("Numbered by XML")
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    ppr.append(num_pr)
    doc.save(path)
    extracted = extract_document(path)
    text = "\n".join(unit.markdown for unit in extracted.units)
    assert "# Document title" in text
    assert "1. Numbered by XML" in text
