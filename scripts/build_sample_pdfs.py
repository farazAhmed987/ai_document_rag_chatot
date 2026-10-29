"""Generate deterministic sample PDFs and document variants.

SRD Sections 8.1, 12 and Appendix A.
Requirements: DM-1, DM-2, IN-1 and IN-2.

Only the five approved Markdown source documents are read.
Original sample data and the internal facts sheet are never changed.

Usage:
    python scripts/build_sample_pdfs.py
    python scripts/build_sample_pdfs.py --company-name "Example Company"
"""

from __future__ import annotations

import argparse
import hashlib
import html
import re
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Pt
from markdown_it import MarkdownIt
from markdown_it.token import Token
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    HRFlowable,
    LongTable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "data" / "sample_docs_src"
OUTPUT_DIR = ROOT / "data" / "sample_docs"

COMPANY_NAME = "Himalaya Crest Realty (Pvt.) Ltd."

DOCUMENTS = (
    "Company_Overview",
    "Property_Listings",
    "Services_and_Fees",
    "FAQs",
    "Policies_and_Terms",
)

PAGE_WIDTH, PAGE_HEIGHT = A4
LEFT_MARGIN = 30
RIGHT_MARGIN = 30
TOP_MARGIN = 55
BOTTOM_MARGIN = 45

CONTENT_WIDTH = PAGE_WIDTH - LEFT_MARGIN - RIGHT_MARGIN


@dataclass
class Block:
    """A semantic Markdown block shared by PDF and DOCX renderers.

    Attributes:
        kind: Heading, paragraph, table, rule, or code.
        inline: Inline Markdown tokens.
        level: Heading level or list nesting depth.
        marker: Optional list marker.
        rows: Table rows, each containing inline-token cells.
        text: Literal code or HTML block content.
    """

    kind: str
    inline: list[Token] = field(default_factory=list)
    level: int = 0
    marker: str = ""
    rows: list[list[list[Token]]] = field(default_factory=list)
    text: str = ""


def validate_winansi(value: str, *, label: str) -> None:
    """Reject characters unsupported by standard PDF fonts.

    Args:
        value: Text to validate.
        label: Name used in the error message.

    Raises:
        ValueError: If the text contains a non-Windows-1252 character.
    """
    try:
        value.encode("cp1252", errors="strict")
    except UnicodeEncodeError as exc:
        character = value[exc.start]
        raise ValueError(
            f"{label} contains U+{ord(character):04X} "
            "which cannot be encoded as Windows-1252"
        ) from exc


def metadata_title(markdown: str) -> str:
    """Read the Document title field from source metadata.

    Args:
        markdown: Complete Markdown source.

    Returns:
        Exact document title, without Markdown emphasis markers.

    Raises:
        ValueError: When required metadata is missing.
    """
    for original_line in markdown.splitlines():
        line = re.sub(r"[*_`]", "", original_line.strip())
        line = re.sub(r"^[-+]\s*", "", line)

        match = re.match(
            r"Document title\s*:\s*(.+)$",
            line,
            flags=re.IGNORECASE,
        )

        if match:
            return match.group(1).strip()

    raise ValueError("Missing required 'Document title:' metadata")


def parse_table(
    tokens: list[Token],
    start: int,
) -> tuple[Block, int]:
    """Parse a Markdown table without flattening its cells.

    Args:
        tokens: Markdown parser token stream.
        start: Index of table_open.

    Returns:
        A table block and index immediately after table_close.
    """
    rows: list[list[list[Token]]] = []
    current_row: list[list[Token]] | None = None

    index = start + 1

    while index < len(tokens):
        token = tokens[index]

        if token.type == "table_close":
            break

        if token.type == "tr_open":
            current_row = []

        elif token.type in ("th_open", "td_open"):
            cell_tokens: list[Token] = []

            if index + 1 < len(tokens):
                following = tokens[index + 1]
                if following.type == "inline":
                    cell_tokens = list(following.children or [])

            if current_row is not None:
                current_row.append(cell_tokens)

        elif token.type == "tr_close":
            if current_row is not None:
                rows.append(current_row)
            current_row = None

        index += 1

    if index >= len(tokens):
        raise ValueError("Unclosed Markdown table")

    if rows:
        width = len(rows[0])

        if any(len(row) != width for row in rows):
            raise ValueError("Inconsistent Markdown table columns")

    return Block(kind="table", rows=rows), index + 1


def parse_markdown(markdown: str) -> list[Block]:
    """Convert Markdown into an ordered stream of semantic blocks.

    Args:
        markdown: Original source content.

    Returns:
        Blocks retaining headings, paragraphs, lists and tables.
    """
    parser = MarkdownIt("commonmark").enable("table")
    tokens = parser.parse(markdown)

    blocks: list[Block] = []
    list_stack: list[dict[str, int]] = []
    pending_marker = ""

    index = 0

    while index < len(tokens):
        token = tokens[index]
        kind = token.type

        if kind == "table_open":
            table, index = parse_table(tokens, index)
            blocks.append(table)
            continue

        if kind in ("bullet_list_open", "ordered_list_open"):
            initial = 0

            if kind == "ordered_list_open":
                initial = int(token.attrGet("start") or 1) - 1

            list_stack.append(
                {
                    "ordered": int(kind == "ordered_list_open"),
                    "counter": initial,
                }
            )

        elif kind in ("bullet_list_close", "ordered_list_close"):
            if list_stack:
                list_stack.pop()

        elif kind == "list_item_open":
            if list_stack:
                current = list_stack[-1]
                current["counter"] += 1

                if current["ordered"]:
                    pending_marker = f"{current['counter']}."
                else:
                    pending_marker = "•"

        elif kind == "list_item_close":
            pending_marker = ""

        elif kind == "heading_open":
            if index + 1 < len(tokens):
                inline = tokens[index + 1]

                blocks.append(
                    Block(
                        kind="heading",
                        inline=list(inline.children or []),
                        level=int(token.tag[1:]),
                    )
                )

        elif kind == "paragraph_open":
            if index + 1 < len(tokens):
                inline = tokens[index + 1]

                blocks.append(
                    Block(
                        kind="paragraph",
                        inline=list(inline.children or []),
                        level=len(list_stack),
                        marker=pending_marker,
                    )
                )

                pending_marker = ""

        elif kind == "hr":
            blocks.append(Block(kind="rule"))

        elif kind in ("fence", "code_block", "html_block"):
            blocks.append(
                Block(
                    kind="code",
                    text=token.content,
                )
            )

        index += 1

    return blocks


def inline_plain(tokens: list[Token]) -> str:
    """Render inline tokens as unformatted text.

    Args:
        tokens: Markdown inline tokens.

    Returns:
        Text with logical whitespace and no formatting markers.
    """
    pieces: list[str] = []

    for token in tokens:
        if token.type in (
            "text",
            "code_inline",
            "html_inline",
            "image",
        ):
            pieces.append(token.content)

        elif token.type in ("softbreak", "hardbreak"):
            pieces.append(" ")

    return "".join(pieces)


def inline_pdf(tokens: list[Token]) -> str:
    """Convert inline Markdown into ReportLab paragraph markup.

    Args:
        tokens: Markdown inline tokens.

    Returns:
        Escaped ReportLab-compatible markup.
    """
    pieces: list[str] = []

    tags = {
        "strong_open": "<b>",
        "strong_close": "</b>",
        "em_open": "<i>",
        "em_close": "</i>",
    }

    for token in tokens:
        kind = token.type

        if kind in tags:
            pieces.append(tags[kind])

        elif kind in ("text", "html_inline", "image"):
            pieces.append(html.escape(token.content, quote=True))

        elif kind == "code_inline":
            pieces.append(
                '<font name="Courier">'
                + html.escape(token.content, quote=True)
                + "</font>"
            )

        elif kind in ("softbreak", "hardbreak"):
            pieces.append(" ")

    return "".join(pieces)


def pdf_styles() -> dict[str, ParagraphStyle]:
    """Build readable, consistent PDF typography.

    Returns:
        Named ReportLab paragraph styles.
    """
    base = ParagraphStyle(
        "Body",
        fontName="Helvetica",
        fontSize=10,
        leading=12.4,
        spaceAfter=5,
        alignment=TA_LEFT,
        allowWidows=0,
        allowOrphans=0,
    )

    return {
        "body": base,
        "h1": ParagraphStyle(
            "H1",
            parent=base,
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=17,
            spaceBefore=10,
            spaceAfter=6,
            keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "H2",
            parent=base,
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            spaceBefore=9,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "h3": ParagraphStyle(
            "H3",
            parent=base,
            fontName="Helvetica-Bold",
            fontSize=10,
            leading=12,
            spaceBefore=7,
            spaceAfter=4,
            keepWithNext=True,
        ),
        "table": ParagraphStyle(
            "TableCell",
            parent=base,
            fontSize=8,
            leading=9.6,
            spaceAfter=0,
            splitLongWords=True,
        ),
    }


def table_widths(rows: list[list[list[Token]]]) -> list[float]:
    """Allocate table width using bounded column-content weights.

    Args:
        rows: Markdown table rows.

    Returns:
        Column widths that sum to the available page width.
    """
    count = len(rows[0])

    weights: list[float] = []

    for column in range(count):
        longest = max(len(inline_plain(row[column])) for row in rows)

        weights.append(float(max(10, min(longest, 35))))

    total = sum(weights)

    return [CONTENT_WIDTH * weight / total for weight in weights]


def make_pdf_story(markdown: str) -> list[object]:
    """Convert Markdown to ReportLab Platypus flowables.

    Args:
        markdown: Source document.

    Returns:
        A flowable sequence with splittable, bordered tables.
    """
    styles = pdf_styles()
    story: list[object] = []

    for block in parse_markdown(markdown):
        if block.kind == "heading":
            level = min(block.level, 3)

            story.append(
                Paragraph(
                    inline_pdf(block.inline),
                    styles[f"h{level}"],
                )
            )

        elif block.kind == "paragraph":
            content = inline_pdf(block.inline)

            if block.marker:
                content = html.escape(block.marker) + "&nbsp;&nbsp;" + content

            paragraph_style = styles["body"]

            if block.level:
                paragraph_style = ParagraphStyle(
                    f"ListDepth{block.level}",
                    parent=styles["body"],
                    leftIndent=12 * block.level,
                )

            story.append(Paragraph(content or "&#160;", paragraph_style))

        elif block.kind == "rule":
            story.append(
                HRFlowable(
                    width="100%",
                    thickness=0.5,
                    color=colors.grey,
                    spaceBefore=6,
                    spaceAfter=6,
                )
            )

        elif block.kind == "code":
            content = html.escape(block.text).replace(
                "\n",
                "<br/>",
            )

            story.append(
                Paragraph(
                    content or "&#160;",
                    styles["body"],
                )
            )

        elif block.kind == "table" and block.rows:
            formatted_rows = []

            for row in block.rows:
                formatted_rows.append(
                    [
                        Paragraph(
                            inline_pdf(cell) or "&#160;",
                            styles["table"],
                        )
                        for cell in row
                    ]
                )

            table = LongTable(
                formatted_rows,
                colWidths=table_widths(block.rows),
                repeatRows=1,
                splitByRow=1,
                hAlign="LEFT",
            )

            table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor("#E7EEF5"),
                        ),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.4,
                            colors.HexColor("#7C8793"),
                        ),
                        (
                            "VALIGN",
                            (0, 0),
                            (-1, -1),
                            "TOP",
                        ),
                        (
                            "LEFTPADDING",
                            (0, 0),
                            (-1, -1),
                            4,
                        ),
                        (
                            "RIGHTPADDING",
                            (0, 0),
                            (-1, -1),
                            4,
                        ),
                        (
                            "TOPPADDING",
                            (0, 0),
                            (-1, -1),
                            3,
                        ),
                        (
                            "BOTTOMPADDING",
                            (0, 0),
                            (-1, -1),
                            3,
                        ),
                    ]
                )
            )

            story.extend(
                [
                    Spacer(1, 5),
                    table,
                    Spacer(1, 6),
                ]
            )

    return story


class NumberedCanvas(canvas.Canvas):
    """Canvas supporting accurate 'Page n of N' numbering.

    SRD Appendix A and Section 8.1.

    Each page state is saved until the total page count is known.
    Invariant PDF mode removes variable creation timestamps and IDs.
    """

    def __init__(self, filename: object, **kwargs: object) -> None:
        """Initialize the deterministic PDF canvas."""
        kwargs.setdefault("invariant", 1)
        super().__init__(filename, **kwargs)

        self._saved_page_states: list[dict[str, object]] = []

    def showPage(self) -> None:
        """Save the current page without finalizing it yet."""
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        """Write every page after the final page count is known."""
        page_count = len(self._saved_page_states)

        for state in self._saved_page_states:
            self.__dict__.update(state)

            self.saveState()
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#374151"))

            self.drawRightString(
                PAGE_WIDTH - RIGHT_MARGIN,
                22,
                f"Page {self._pageNumber} of {page_count}",
            )

            self.restoreState()
            super().showPage()

        super().save()


def build_pdf_bytes(
    markdown: str,
    *,
    company_name: str,
    filename: str,
) -> bytes:
    """Render one complete deterministic A4 PDF.

    Args:
        markdown: Source Markdown.
        company_name: Header company name.
        filename: Source name for encoding diagnostics.

    Returns:
        Fully generated PDF bytes.
    """
    validate_winansi(markdown, label=filename)
    validate_winansi(company_name, label="Company name")

    title = metadata_title(markdown)
    validate_winansi(title, label="Document title")

    buffer = BytesIO()

    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=LEFT_MARGIN,
        rightMargin=RIGHT_MARGIN,
        topMargin=TOP_MARGIN,
        bottomMargin=BOTTOM_MARGIN,
        title=title,
        author=company_name,
        subject="Fictional sample document for RAG testing",
        pageCompression=1,
        allowSplitting=1,
        invariant=1,
    )

    def draw_header(
        drawing_canvas: canvas.Canvas,
        document: SimpleDocTemplate,
    ) -> None:
        """Draw the running header on each page."""
        del document

        drawing_canvas.saveState()
        drawing_canvas.setFont("Helvetica", 8)

        drawing_canvas.drawString(
            LEFT_MARGIN,
            PAGE_HEIGHT - 26,
            f"{company_name} | {title}",
        )

        drawing_canvas.setStrokeColor(colors.HexColor("#9CA3AF"))

        drawing_canvas.line(
            LEFT_MARGIN,
            PAGE_HEIGHT - 32,
            PAGE_WIDTH - RIGHT_MARGIN,
            PAGE_HEIGHT - 32,
        )

        drawing_canvas.restoreState()

    pdf.build(
        make_pdf_story(markdown),
        onFirstPage=draw_header,
        onLaterPages=draw_header,
        canvasmaker=NumberedCanvas,
    )

    return buffer.getvalue()


def append_docx_inline(
    paragraph: object,
    tokens: list[Token],
) -> None:
    """Write Markdown inline formatting into DOCX paragraph runs.

    Args:
        paragraph: python-docx paragraph.
        tokens: Inline Markdown tokens.
    """
    bold_depth = 0
    italic_depth = 0

    for token in tokens:
        kind = token.type

        if kind == "strong_open":
            bold_depth += 1
            continue

        if kind == "strong_close":
            bold_depth = max(0, bold_depth - 1)
            continue

        if kind == "em_open":
            italic_depth += 1
            continue

        if kind == "em_close":
            italic_depth = max(0, italic_depth - 1)
            continue

        if kind in (
            "text",
            "code_inline",
            "html_inline",
            "image",
        ):
            value = token.content

        elif kind in ("softbreak", "hardbreak"):
            value = " "

        else:
            continue

        run = paragraph.add_run(value)
        run.bold = bold_depth > 0
        run.italic = italic_depth > 0

        if kind == "code_inline":
            run.font.name = "Courier New"


def build_docx(markdown: str, target: Path) -> None:
    """Create a real, structured DOCX document from Markdown.

    Args:
        markdown: Source Markdown.
        target: Destination file.

    The output retains Heading styles, lists, paragraphs,
    inline emphasis and editable Word tables.
    """
    document = Document()

    document.core_properties.title = metadata_title(markdown)
    document.core_properties.author = COMPANY_NAME
    document.core_properties.created = datetime(2000, 1, 1)
    document.core_properties.modified = datetime(2000, 1, 1)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    for block in parse_markdown(markdown):
        if block.kind == "heading":
            paragraph = document.add_heading(level=min(block.level, 9))

            append_docx_inline(paragraph, block.inline)

        elif block.kind == "paragraph":
            style = None

            if block.marker:
                style = (
                    "List Number"
                    if block.marker.endswith(".") and block.marker[:-1].isdigit()
                    else "List Bullet"
                )

            paragraph = document.add_paragraph(style=style)

            if block.level > 1:
                paragraph.paragraph_format.left_indent = Pt(18 * (block.level - 1))

            append_docx_inline(paragraph, block.inline)

        elif block.kind == "table" and block.rows:
            table = document.add_table(
                rows=0,
                cols=len(block.rows[0]),
            )

            table.style = "Table Grid"

            for row_index, row in enumerate(block.rows):
                cells = table.add_row().cells

                for cell_index, token_list in enumerate(row):
                    paragraph = cells[cell_index].paragraphs[0]
                    append_docx_inline(paragraph, token_list)

                    if row_index == 0:
                        for run in paragraph.runs:
                            run.bold = True

        elif block.kind == "rule":
            document.add_paragraph("─" * 30)

        elif block.kind == "code":
            document.add_paragraph(block.text)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


def build_documents(
    *,
    company_name: str,
    source_dir: Path,
    output_dir: Path,
) -> dict[str, str]:
    """Generate the five PDFs and three source-format variants.

    Args:
        company_name: Running header company name.
        source_dir: Directory containing the approved Markdown files.
        output_dir: Destination for generated documents.

    Returns:
        Mapping of generated PDF filename to SHA-256 digest.
    """
    validate_winansi(company_name, label="Company name")

    missing = [stem for stem in DOCUMENTS if not (source_dir / f"{stem}.md").is_file()]

    if missing:
        raise FileNotFoundError(
            "Missing required source document(s): " + ", ".join(missing)
        )

    output_dir.mkdir(parents=True, exist_ok=True)

    variants_dir = output_dir / "variants"
    variants_dir.mkdir(parents=True, exist_ok=True)

    checksums: dict[str, str] = {}

    for stem in DOCUMENTS:
        source = source_dir / f"{stem}.md"
        markdown = source.read_text(encoding="utf-8")

        pdf_bytes = build_pdf_bytes(
            markdown,
            company_name=company_name,
            filename=source.name,
        )

        destination = output_dir / f"{stem}.pdf"
        destination.write_bytes(pdf_bytes)

        checksums[destination.name] = hashlib.sha256(pdf_bytes).hexdigest()

    services_markdown = (source_dir / "Services_and_Fees.md").read_text(
        encoding="utf-8"
    )

    build_docx(
        services_markdown,
        variants_dir / "Services_and_Fees.docx",
    )

    faq_source = source_dir / "FAQs.md"
    faq_bytes = faq_source.read_bytes()

    (variants_dir / "FAQs.txt").write_text(
        faq_bytes.decode("utf-8"),
        encoding="utf-8",
        newline="\n",
    )

    shutil.copyfile(
        faq_source,
        variants_dir / "FAQs.md",
    )

    return checksums


def main() -> None:
    """Parse CLI arguments and generate all Phase 1 artifacts."""
    parser = argparse.ArgumentParser(
        description="Build five deterministic sample PDFs."
    )

    parser.add_argument(
        "--company-name",
        default=COMPANY_NAME,
    )

    parser.add_argument(
        "--source-dir",
        type=Path,
        default=SOURCE_DIR,
    )

    parser.add_argument(
        "--output-dir",
        type=Path,
        default=OUTPUT_DIR,
    )

    arguments = parser.parse_args()

    checksums = build_documents(
        company_name=arguments.company_name,
        source_dir=arguments.source_dir,
        output_dir=arguments.output_dir,
    )

    print("Generated sample PDFs:")
    for filename, digest in checksums.items():
        print(f"{filename}: {digest}")

    print("Generated variants:")
    print("  variants/Services_and_Fees.docx")
    print("  variants/FAQs.txt")
    print("  variants/FAQs.md")


if __name__ == "__main__":
    main()
